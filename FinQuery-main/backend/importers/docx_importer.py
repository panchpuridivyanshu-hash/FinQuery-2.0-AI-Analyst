from docx import Document


def read_docx_file(file_path):
    document = Document(file_path)

    text = ""

    # Read normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Read tables
    for table in document.tables:
        for row in table.rows:
            row_data = []

            for cell in row.cells:
                row_data.append(cell.text.strip())

            text += " ".join(row_data) + "\n"

    return text

import re


def parse_docx_transactions(text):
    transactions = []

    pattern = re.compile(
        r"Date:\s*(.*?)\s*"
        r"Description:\s*(.*?)\s*"
        r"Category:\s*(.*?)\s*"
        r"Amount:\s*([\d,]+(?:\.\d+)?)",
        re.IGNORECASE
    )

    matches = pattern.findall(text)

    for match in matches:

        transaction = {
            "date": match[0].strip(),
            "description": match[1].strip(),
            "category": match[2].strip(),
            "amount": match[3].replace(",", "").strip()
        }

        transactions.append(transaction)

    if transactions:
        return transactions

    known_categories = [
        "Food & Groceries",
        "Food & Dining",
        "Entertainment",
        "Healthcare",
        "Education",
        "Utilities",
        "Transport",
        "Shopping",
        "Travel",
        "Bills",
    ]

    table_pattern = re.compile(
        r"^(?:[A-Z]{2,}\d+\s+)?"
        r"(\d{2}[-/]\d{2}[-/]\d{4})\s+"
        r"(.+?)\s+"
        r"(?:Expense|Income)\s+"
        r"([\d,]+(?:\.\d+)?)"
        r"(?:\s+.*)?$",
        re.IGNORECASE,
    )

    for line in text.splitlines():
        match = table_pattern.match(" ".join(line.split()))
        if not match:
            continue

        date_value, details, amount = match.groups()
        category = next(
            (item for item in known_categories if details.lower().endswith(item.lower())),
            "Uncategorized",
        )
        description = details[: -len(category)].strip() if category != "Uncategorized" else details

        transactions.append({
            "date": date_value,
            "description": description,
            "category": category,
            "amount": amount.replace(",", ""),
        })

    return transactions
